from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


DOCX_PATH = Path(__file__).resolve().parents[1] / "docs" / "Porjai_Abstract_HIS_API_v5_draft.docx"


REPLACEMENTS = {
    '"clinic_name": "Kidney Clinic",': '"clinic_name": "ศูนย์ไตเทียม 1",',
    '"coverage_code": "OFC",': '"pttype": "10",',
    '"coverage_name": "Out-of-pocket"': '"pttype_name": "ชำระเงินสด ทั่วไป"',
}


def insert_paragraph_after(paragraph, text: str):
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    for candidate in paragraph._parent.paragraphs:
        if candidate._element is new_p:
            candidate.text = text
            return candidate
    raise RuntimeError("Inserted paragraph could not be found")


def main() -> None:
    doc = Document(DOCX_PATH)

    # API 1 response table was edited after the meeting; clean examples so
    # they remain single-value examples and match the new field names.
    response_fields = doc.tables[2]
    for row in response_fields.rows:
        field = row.cells[0].text.strip()
        if field == "clinic_code":
            row.cells[3].text = "1004"
            row.cells[4].text = "Values: 1004, 1005"
        elif field == "clinic_name":
            row.cells[3].text = "ศูนย์ไตเทียม 1"
            row.cells[4].text = ""

    # Keep API 2 clinic scope consistent with the updated table.
    request_fields = doc.tables[3]
    for row in request_fields.rows:
        if row.cells[0].text.strip() == "clinic_code":
            row.cells[4].text = "Values: 1004, 1005"

    for paragraph in list(doc.paragraphs):
        text = paragraph.text
        for old, new in REPLACEMENTS.items():
            if old in text:
                text = text.replace(old, new)
        paragraph.text = text

    # The order-item table contains med_ned; include it in the examples too.
    paragraphs = list(doc.paragraphs)
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        next_text = paragraphs[idx + 1].text.strip() if idx + 1 < len(paragraphs) else ""
        if text == '"meduestype": "IV1"' and next_text != '"med_ned": 8':
            paragraph.text = '"meduestype": "IV1",'
            insert_paragraph_after(paragraph, '"med_ned": 8')

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    main()
