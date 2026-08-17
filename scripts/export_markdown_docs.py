from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt


INLINE_CODE_RE = re.compile(r"`([^`]+)`")
ORDERED_RE = re.compile(r"^(\d+)\.\s+(.*)$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
CODE_FENCE_RE = re.compile(r"^```")


def split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
      text = text[1:]
    if text.endswith("|"):
      text = text[:-1]
    return [cell.strip() for cell in text.split("|")]


def is_table_delimiter(line: str) -> bool:
    cells = split_table_row(line)
    if not cells:
        return False
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def ensure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    if "CodeInline" not in doc.styles:
        code_style = doc.styles.add_style("CodeInline", WD_STYLE_TYPE.CHARACTER)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9.5)


def add_inline_runs(paragraph, text: str) -> None:
    if not text:
        return
    last = 0
    for match in INLINE_CODE_RE.finditer(text):
        if match.start() > last:
            paragraph.add_run(text[last:match.start()])
        run = paragraph.add_run(match.group(1))
        run.style = "CodeInline"
        last = match.end()
    if last < len(text):
        paragraph.add_run(text[last:])


def finalize_paragraph(doc: Document, lines: list[str]) -> None:
    text = " ".join(line.strip() for line in lines if line.strip()).strip()
    if not text:
        return
    paragraph = doc.add_paragraph()
    add_inline_runs(paragraph, text)


def add_heading(doc: Document, level: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {min(max(level, 1), 4)}"
    add_inline_runs(paragraph, text.strip())


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    add_inline_runs(paragraph, text.strip())


def add_ordered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    add_inline_runs(paragraph, text.strip())


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_index, row_values in enumerate(rows):
        row = table.add_row().cells
        for cell_index in range(col_count):
            value = row_values[cell_index] if cell_index < len(row_values) else ""
            paragraph = row[cell_index].paragraphs[0]
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
            add_inline_runs(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line or " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)


def convert_markdown_file(source: Path, destination: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    ensure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.page_width = doc.sections[0].page_width
    section.page_height = doc.sections[0].page_height

    paragraph_buffer: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            finalize_paragraph(doc, paragraph_buffer)
            paragraph_buffer = []
            i += 1
            continue

        if CODE_FENCE_RE.match(stripped):
            finalize_paragraph(doc, paragraph_buffer)
            paragraph_buffer = []
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not CODE_FENCE_RE.match(lines[i].strip()):
                code_lines.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, code_lines)
            if i < len(lines):
                i += 1
            continue

        heading_match = HEADING_RE.match(stripped)
        ordered_match = ORDERED_RE.match(stripped)
        table_line = stripped.startswith("|") and "|" in stripped[1:]

        if heading_match or stripped.startswith("- ") or ordered_match or table_line:
            finalize_paragraph(doc, paragraph_buffer)
            paragraph_buffer = []

        if heading_match:
            add_heading(doc, len(heading_match.group(1)), heading_match.group(2))
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        if ordered_match:
            add_ordered(doc, ordered_match.group(2))
            i += 1
            continue

        if table_line:
            table_rows: list[list[str]] = []
            while i < len(lines):
                current = lines[i].strip()
                if not current.startswith("|") or "|" not in current[1:]:
                    break
                if is_table_delimiter(current):
                    i += 1
                    continue
                table_rows.append(split_table_row(current))
                i += 1
            add_table(doc, table_rows)
            continue

        paragraph_buffer.append(stripped)
        i += 1

    finalize_paragraph(doc, paragraph_buffer)
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export markdown documents to DOCX.")
    parser.add_argument("inputs", nargs="+", help="Markdown file(s) to convert")
    parser.add_argument("--output-dir", default="generated-docs", help="Directory for output DOCX files")
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    for raw_input in args.inputs:
        source = Path(raw_input)
        if not source.exists():
            raise FileNotFoundError(f"Input not found: {source}")
        destination = output_dir / f"{source.stem}.docx"
        convert_markdown_file(source, destination)
        print(f"[DOCX] {source} -> {destination}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
